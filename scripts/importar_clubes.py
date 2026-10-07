#!/usr/bin/env python3
"""Converte o Excel dos clubes (pt/clubes_redes_sociais.xlsx) no pt/clubes.json que o servidor lê.

Junta três coisas:
  - a folha «Redes sociais» (Instagram, Facebook e site de cada clube);
  - a folha «Clubes 2026-27» (em que divisão e série joga cada equipa);
  - o pt/redes-encontradas.json (links encontrados depois, por pesquisa ou pelo `npm run procurar-redes`),
    que só preenche o que no Excel está vazio.

Uso:  python3 scripts/importar_clubes.py [caminho/do/excel.xlsx]
Precisa do openpyxl (pip install openpyxl).
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
EXCEL = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "pt" / "clubes_redes_sociais.xlsx"
ENCONTRADAS = RAIZ / "pt" / "redes-encontradas.json"
SAIDA = RAIZ / "pt" / "clubes.json"


def limpa(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def handle_ig(url):
    """https://www.instagram.com/clube/ → clube"""
    if not url:
        return None
    m = re.search(r"instagram\.com/([A-Za-z0-9_.]+)", url)
    if not m or m.group(1).lower() in {"p", "reel", "explore", "stories", "accounts"}:
        return None
    return m.group(1).lower()


def link_fb(url):
    if not url or "facebook.com" not in url:
        return None
    url = url.strip().split("#")[0]
    if "profile.php" in url:
        m = re.search(r"id=(\d+)", url)
        return f"https://www.facebook.com/profile.php?id={m.group(1)}" if m else None
    return url.split("?")[0]


def lista(v):
    if not v:
        return []
    return [x.strip() for x in re.split(r"[;,|]", str(v)) if x.strip()]


def main():
    wb = openpyxl.load_workbook(EXCEL, data_only=True)

    # participações: equipa → divisões/séries
    part = {}
    for r in wb["Clubes 2026-27"].iter_rows(min_row=2, values_only=True):
        if not r[7]:
            continue
        chave = (limpa(r[3]), limpa(r[7]))
        part.setdefault(chave, []).append({
            "ambito": limpa(r[1]), "nivelNacional": r[2], "nivelAssoc": r[4],
            "divisao": limpa(r[5]), "serie": limpa(r[6]), "tipo": limpa(r[8]), "estado": limpa(r[9]),
        })

    encontradas = {}
    if ENCONTRADAS.exists():
        for c in json.loads(ENCONTRADAS.read_text("utf-8")).get("clubes", []):
            encontradas[(c.get("assoc"), c["clube"])] = c

    clubes = []
    for r in wb["Redes sociais"].iter_rows(min_row=2, values_only=True):
        if r[0] != "Clube" or not r[1]:
            continue
        nome, assoc = limpa(r[1]), limpa(r[2])
        ig, fb = handle_ig(limpa(r[6])), link_fb(limpa(r[7]))
        est_ig, est_fb = limpa(r[8]), limpa(r[9])
        extra = encontradas.get((assoc, nome))
        if extra:
            if not ig and extra.get("ig"):
                ig, est_ig = handle_ig(extra["ig"]), f"Pesquisa ({extra.get('conf_ig') or 'media'})"
            if not fb and extra.get("fb"):
                fb, est_fb = link_fb(extra["fb"]), f"Pesquisa ({extra.get('conf_fb') or 'media'})"
        variantes = lista(r[5]) or [nome]
        comps = []
        for v in variantes:
            for a in lista(assoc) or [assoc]:
                comps += part.get((a, v), [])
        clubes.append({
            "nome": nome,
            "assoc": assoc,
            "variantes": variantes,
            "divisoes": lista(r[3]),
            "competicoes": comps,
            "instagram": ig,
            "facebook": fb,
            "estadoInstagram": est_ig,
            "estadoFacebook": est_fb,
            "site": limpa(r[14]),
            "pesquisaInstagram": limpa(r[10]),
            "pesquisaFacebook": limpa(r[11]),
        })

    associacoes = []
    for r in wb["Associações - links"].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        res = limpa(r[6]) or ""
        m = re.search(r"associationId=(\d+)", res)
        associacoes.append({
            "nome": limpa(r[0]), "instagram": handle_ig(limpa(r[1])), "facebook": link_fb(limpa(r[2])),
            "site": limpa(r[4]), "resultados": res or None, "fpfId": int(m.group(1)) if m else None,
        })

    SAIDA.write_text(json.dumps({"origem": EXCEL.name, "associacoes": associacoes, "clubes": clubes}, ensure_ascii=False, indent=1), "utf-8")
    com_ig = sum(1 for c in clubes if c["instagram"])
    com_fb = sum(1 for c in clubes if c["facebook"])
    print(f"{len(clubes)} clubes ({com_ig} com Instagram, {com_fb} com Facebook) e {len(associacoes)} associações → {SAIDA}")


if __name__ == "__main__":
    main()
